import re
from collections import Counter, defaultdict

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.housekeeping.models import HKArea, HKWorker


AREA_ALIASES = {
    'aibaan': 'Aibaan',
    'beauki': 'Beuaki',
    'block-a, jrha': 'Jeet Hostel Block A',
    'block-b, jrha': 'Jeet Hostel Block B',
    'block-c, jrha': 'Jeet Hostel Block C',
    'block-d, jrha': 'Jeet Hostel Block D',
    'chimair': 'Chimair',
    'common area & mess': 'Hostel Common Area',
    'duven': 'Duven',
    'emiet': 'Emiet',
    'firpeal': 'Firpeal',
    'griwiksh': 'Griwiksh',
    'hiqom': 'Hiqom',
    'ijokha': 'Ijokha',
    'jurqia': 'Jurqia',
    'kyzeel': 'Kyzeel',
    'lekhaag': 'Lekhaag',
}

SUPERVISOR_ALIASES = {
    'ms. priyanka': 'priyanka.parmar',
    'ms. anita vaghela': 'anita.vaghela',
    'ms. manisha thapa': 'manisha.thapa',
    'mr. vishwajeet mishra': 'vishwajeet.mishra',
    'mr. pankaj a bihola': 'pankaj.bihola',
    'mr. anilkumar solanki': 'anil.solanki',
    'mr. nikulsinh bihola': 'nikul.bihola',
    'ms. rushita chauhan': 'rhushita.chavda',
    'mr. jaimin parmar': 'jaimin.parmar',
    'mr. ramsingh': 'ramprakash.parmar',
}

HEADER_ALIASES = {
    'pay_code': {'cod', 'code', 'pay code', 'employee code', 'emp code'},
    'name': {'name', 'staff name', 'worker name'},
    'category': {'caretogery', 'category', 'worker category'},
    'area': {'hostel', 'area', 'area assignment', 'hostel / area assignment'},
    'supervisor': {'name of supervisor', 'supervisor', 'supervisor name'},
    'work': {'work', 'work allocation', 'allocation'},
}


def clean_text(value):
    if value is None:
        return ''
    text = str(value).replace('\xa0', ' ').strip()
    return re.sub(r'\s+', ' ', text)


def norm_key(value):
    return clean_text(value).casefold()


def norm_name(value):
    return re.sub(r'[^a-z0-9]+', ' ', clean_text(value).casefold()).strip()


def norm_pay_code(value):
    if value is None:
        return ''
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).rstrip('0').rstrip('.')
    text = clean_text(value)
    if not text or text.casefold() in {'nan', 'none', 'null'}:
        return ''
    if re.fullmatch(r'\d+\.0+', text):
        return text.split('.', 1)[0]
    return text


def titleish(value):
    text = clean_text(value)
    if not text:
        return ''
    lower_words = {'ben', 'bhai', 'm', 'k', 'j'}
    words = []
    for word in text.split():
        if word.casefold() in lower_words:
            words.append(word.casefold())
        else:
            words.append(word[:1].upper() + word[1:].lower())
    return ' '.join(words)


def infer_worker_type(category, work):
    category_key = norm_key(category)
    work_key = norm_key(work)
    if category_key == 'reliver' or 'reliever' in category_key:
        return 'reliever', True
    if 'washroom' in work_key or 'waahroom' in work_key:
        return 'washroom', False
    if 'room' in work_key and 'common room' not in work_key:
        return 'room', False
    if any(token in work_key for token in ['floor', 'staircase', 'balcany', 'balcony', 'dusting']):
        return 'floor', False
    return 'general', False


def area_type_for_name(name):
    lowered = norm_key(name)
    if lowered.startswith('jeet hostel') or lowered in {
        'aibaan', 'beuaki', 'chimair', 'duven', 'emiet', 'firpeal', 'griwiksh',
        'hiqom', 'ijokha', 'jurqia', 'kyzeel', 'lekhaag', 'hostel common area',
    }:
        return 'hostel'
    return 'other'


class Command(BaseCommand):
    help = 'Import latest housekeeping staff Excel as idempotent HKWorker upserts.'

    def add_arguments(self, parser):
        parser.add_argument('xlsx_path', help='Path to the Housekeeping staff .xlsx file.')
        parser.add_argument('--dry-run', action='store_true', help='Parse and report without saving changes.')

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError('openpyxl is required. Install backend requirements first.') from exc

        xlsx_path = options['xlsx_path']
        dry_run = options['dry_run']

        workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
        sheet = workbook['all name  '] if 'all name  ' in workbook.sheetnames else workbook.active
        rows = self._parse_sheet(sheet)
        duplicate_pay_codes = {
            code: row_numbers
            for code, row_numbers in self._duplicate_pay_codes(rows).items()
            if len(row_numbers) > 1
        }
        duplicate_rows = {row for row_numbers in duplicate_pay_codes.values() for row in row_numbers[1:]}

        stats = Counter()
        details = defaultdict(list)

        context = self._build_context(details)

        with transaction.atomic():
            for row in rows:
                if row['row'] in duplicate_rows:
                    stats['duplicate_rows'] += 1
                    continue
                try:
                    result = self._upsert_worker(row, context, details, dry_run)
                    stats[result] += 1
                except Exception as exc:
                    stats['errors'] += 1
                    details['Errors'].append(f"Row {row['row']}: {exc}")
            if dry_run:
                transaction.set_rollback(True)

        self._print_report(rows, stats, details, duplicate_pay_codes, dry_run)

    def _parse_sheet(self, sheet):
        raw_rows = list(sheet.iter_rows(values_only=True))
        if not raw_rows:
            raise CommandError('Workbook has no rows.')

        headers = [norm_key(cell) for cell in raw_rows[0]]
        col_map = {}
        for field, aliases in HEADER_ALIASES.items():
            for idx, header in enumerate(headers):
                if header in aliases:
                    col_map[field] = idx
                    break
            if field not in col_map:
                raise CommandError(f'Missing required column for {field}. Headers: {headers}')

        rows = []
        for row_num, values in enumerate(raw_rows[1:], start=2):
            row = {
                'row': row_num,
                'pay_code': norm_pay_code(values[col_map['pay_code']] if len(values) > col_map['pay_code'] else ''),
                'name': clean_text(values[col_map['name']] if len(values) > col_map['name'] else ''),
                'category': clean_text(values[col_map['category']] if len(values) > col_map['category'] else ''),
                'area_raw': clean_text(values[col_map['area']] if len(values) > col_map['area'] else ''),
                'supervisor_raw': clean_text(values[col_map['supervisor']] if len(values) > col_map['supervisor'] else ''),
                'work': clean_text(values[col_map['work']] if len(values) > col_map['work'] else ''),
            }
            if not any(row[key] for key in ['pay_code', 'name', 'category', 'area_raw', 'supervisor_raw', 'work']):
                continue
            rows.append(row)
        return rows

    def _duplicate_pay_codes(self, rows):
        pay_rows = defaultdict(list)
        for row in rows:
            if row['pay_code']:
                pay_rows[row['pay_code']].append(row['row'])
        return pay_rows

    def _build_context(self, details):
        User = get_user_model()
        areas = {norm_key(area.name): area for area in HKArea.objects.all()}
        supervisors = {}
        for label, username in SUPERVISOR_ALIASES.items():
            user = User.objects.filter(username=username, role__in=['supervisor', 'caretaker']).first()
            if user:
                supervisors[label] = user
            else:
                details['Supervisor Not Found'].append(f'{label} -> {username}')

        workers = list(HKWorker.objects.select_related('area', 'supervisor').all())
        workers_by_pay = defaultdict(list)
        workers_by_name_area = defaultdict(list)
        for worker in workers:
            pay_code = norm_pay_code(worker.pay_code)
            if pay_code:
                workers_by_pay[pay_code].append(worker)
            if worker.area:
                workers_by_name_area[(norm_name(worker.name), norm_key(worker.area.name))].append(worker)

        return {
            'areas': areas,
            'supervisors': supervisors,
            'workers_by_pay': workers_by_pay,
            'workers_by_name_area': workers_by_name_area,
        }

    def _resolve_area(self, row, context, details, dry_run):
        target_name = AREA_ALIASES.get(norm_key(row['area_raw']), clean_text(row['area_raw']))
        if not target_name:
            return None
        area = context['areas'].get(norm_key(target_name))
        if area:
            return area
        if dry_run:
            details['Areas Created'].append(target_name)
            return HKArea(name=target_name, area_type=area_type_for_name(target_name), is_active=True)
        area = HKArea.objects.create(name=target_name, area_type=area_type_for_name(target_name), is_active=True)
        context['areas'][norm_key(area.name)] = area
        details['Areas Created'].append(area.name)
        return area

    def _resolve_supervisor(self, row, context):
        return context['supervisors'].get(norm_key(row['supervisor_raw']))

    def _match_worker(self, row, area, context, details):
        if row['pay_code']:
            matches = context['workers_by_pay'].get(row['pay_code'], [])
            if len(matches) == 1:
                return matches[0]
            if len(matches) > 1:
                details['Ambiguous Matches'].append(f"Row {row['row']}: pay code {row['pay_code']} matched multiple workers")
                return 'ambiguous'
        if area:
            matches = context['workers_by_name_area'].get((norm_name(row['name']), norm_key(area.name)), [])
            if len(matches) == 1:
                return matches[0]
            if len(matches) > 1:
                details['Ambiguous Matches'].append(f"Row {row['row']}: name + area matched multiple workers")
                return 'ambiguous'
        return None

    def _upsert_worker(self, row, context, details, dry_run):
        if not row['name']:
            details['Errors'].append(f"Row {row['row']}: missing worker name")
            return 'errors'

        area = self._resolve_area(row, context, details, dry_run)
        supervisor = self._resolve_supervisor(row, context)
        worker_type, is_reliever = infer_worker_type(row['category'], row['work'])
        worker = self._match_worker(row, area, context, details)
        if worker == 'ambiguous':
            return 'ambiguous_matches'

        name = titleish(row['name'])
        if worker is None:
            if not dry_run:
                worker = HKWorker.objects.create(
                    name=name,
                    pay_code=row['pay_code'],
                    area=area if getattr(area, 'pk', None) else None,
                    supervisor=supervisor,
                    worker_type=worker_type,
                    shift='day',
                    is_active=True,
                    is_reliever=is_reliever,
                )
                if row['pay_code']:
                    context['workers_by_pay'][row['pay_code']].append(worker)
                if worker.area:
                    context['workers_by_name_area'][(norm_name(worker.name), norm_key(worker.area.name))].append(worker)
            return 'workers_created'

        updates = []
        update_values = {}
        if name and worker.name != name:
            update_values['name'] = name
        if row['pay_code'] and norm_pay_code(worker.pay_code) != row['pay_code']:
            update_values['pay_code'] = row['pay_code']
        if area and getattr(area, 'pk', None) and worker.area_id != area.pk:
            update_values['area'] = area
        if supervisor and worker.supervisor_id != supervisor.pk:
            update_values['supervisor'] = supervisor
        if worker_type and worker.worker_type != worker_type:
            update_values['worker_type'] = worker_type
        if worker.is_reliever != is_reliever:
            update_values['is_reliever'] = is_reliever

        if not update_values:
            return 'workers_skipped'

        if not dry_run:
            for field, value in update_values.items():
                setattr(worker, field, value)
                updates.append(field)
            worker.save(update_fields=updates)
        return 'workers_updated'

    def _print_report(self, rows, stats, details, duplicate_pay_codes, dry_run):
        self.stdout.write(self.style.WARNING('DRY RUN - no changes saved') if dry_run else self.style.SUCCESS('Housekeeping Staff Import Complete'))
        self.stdout.write('')
        self.stdout.write(f"Rows Parsed: {len(rows)}")
        self.stdout.write(f"Workers Created: {stats['workers_created']}")
        self.stdout.write(f"Workers Updated: {stats['workers_updated']}")
        self.stdout.write(f"Workers Skipped: {stats['workers_skipped']}")
        self.stdout.write('Users Created: 0')
        self.stdout.write('Users Updated: 0')
        self.stdout.write(f"Areas Created: {len(details['Areas Created'])}")
        self.stdout.write(f"Supervisor Not Found: {len(details['Supervisor Not Found'])}")
        self.stdout.write(f"Duplicate Rows: {stats['duplicate_rows']}")
        self.stdout.write(f"Ambiguous Matches: {stats['ambiguous_matches']}")
        self.stdout.write(f"Errors: {stats['errors']}")

        if duplicate_pay_codes:
            self.stdout.write('\nDuplicate Records:')
            for pay_code, row_numbers in duplicate_pay_codes.items():
                self.stdout.write(f"- Pay Code {pay_code} appears in rows {', '.join(map(str, row_numbers))}")

        for section in ['Supervisor Not Found', 'Areas Created', 'Ambiguous Matches', 'Errors']:
            items = details.get(section) or []
            if not items:
                continue
            self.stdout.write(f'\n{section}:')
            for item in items:
                self.stdout.write(f'- {item}')
