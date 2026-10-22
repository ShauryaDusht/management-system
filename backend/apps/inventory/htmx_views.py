import csv
import io
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.contrib.auth import authenticate, login as auth_login, logout as auth_logout, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import F, Prefetch, Sum, Value, DecimalField
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST

from apps.users.models import normalize_username_value
from .models import Category, Item, ItemIssue, ItemUsageLog, ItemRequest, ItemRequestBatch, PurchaseRequest, RequisitionItem, StockEntry
from .request_permissions import (
    EXCESSIVE_QUANTITY_MESSAGE,
    UNAUTHORIZED_ITEM_REQUEST_MESSAGE,
    can_request_assets,
    filter_requestable_items,
    user_can_request_item,
)


# ─────────────────────────────────────────────────
# AUTH VIEWS
# ─────────────────────────────────────────────────

def mobile_login(request):
    if request.user.is_authenticated:
        if getattr(request.user, 'role', '') in ['admin', 'inventory_manager']:
            return redirect('/inventory/mobile/')
        return redirect('/portal/login/')

    allowed_roles = ['admin', 'inventory_manager']
    error = ''
    next_url = '/inventory/mobile/'

    if request.method == 'POST':
        username = normalize_username_value(request.POST.get('username', ''))
        password = request.POST.get('password', '')
        raw_next = request.POST.get('next', '').strip()

        user = authenticate(request, username=username, password=password)
        if user is not None and user.is_active and getattr(user, 'role', '') in allowed_roles:
            auth_login(request, user)
            if raw_next and url_has_allowed_host_and_scheme(
                url=raw_next,
                allowed_hosts={request.get_host()},
                require_https=request.is_secure(),
            ):
                return redirect(raw_next)
            return redirect('/inventory/mobile/')
        error = 'You do not have inventory portal access.' if user and user.is_active else 'Invalid username or password.'
    else:
        raw_next = request.GET.get('next', '').strip()
        if raw_next and url_has_allowed_host_and_scheme(
            url=raw_next,
            allowed_hosts={request.get_host()},
            require_https=request.is_secure(),
        ):
            next_url = raw_next

    return render(request, 'inventory/mobile/login.html', {
        'error': error,
        'next': next_url,
    })


@require_POST
def mobile_logout(request):
    auth_logout(request)
    return redirect('/portal/login/')


# ─────────────────────────────────────────────────
# DASHBOARD
# ─────────────────────────────────────────────────

@login_required
def mobile_dashboard(request):
    items = Item.objects.all()
    low_stock = items.filter(
        available_quantity__gt=0,
        available_quantity__lte=F('threshold_quantity'),
    )
    pending = ItemRequest.objects.filter(status='pending')

    context = {
        'total_items': items.count(),
        'low_stock_items': low_stock[:3],
        'low_stock_count': low_stock.count(),
        'out_of_stock_count': items.filter(available_quantity=0).count(),
        'pending_requests': pending.select_related('requested_by', 'item')[:10],
        'pending_count': pending.count(),
        'role': getattr(request.user, 'role', ''),
    }
    return render(request, 'inventory/mobile/dashboard.html', context)


# ─────────────────────────────────────────────────
# INVENTORY LIST  (paginated + search + low_stock filter)
# ─────────────────────────────────────────────────

@login_required
def mobile_inventory_list(request):
    category_id = request.GET.get('category', '')
    search_query = request.GET.get('search', '').strip()
    low_stock_filter = request.GET.get('low_stock') == 'true'
    item_type_filter = (request.GET.get('item_type') or 'all').strip().lower()
    if item_type_filter not in {'all', 'asset', 'consumable'}:
        item_type_filter = 'all'

    base_items_qs = filter_requestable_items(Item.objects.select_related('category').all(), request.user)
    items_qs = base_items_qs
    categories = Category.objects.filter(items__in=base_items_qs).distinct().order_by('name')

    if category_id:
        items_qs = items_qs.filter(category_id=category_id)

    if search_query:
        items_qs = items_qs.filter(name__icontains=search_query)

    if item_type_filter in {'asset', 'consumable'}:
        items_qs = items_qs.filter(item_type=item_type_filter)

    if low_stock_filter:
        items_qs = items_qs.filter(
            available_quantity__gt=0,
            available_quantity__lte=F('threshold_quantity'),
        )

    # Pagination — 15 items per page
    paginator = Paginator(items_qs, 15)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'inventory/mobile/list.html', {
        'items': page_obj,
        'page_obj': page_obj,
        'categories': categories,
        'selected_category': category_id,
        'total_count': base_items_qs.count(),
        'search_query': search_query,
        'low_stock_filter': low_stock_filter,
        'selected_item_type': item_type_filter,
        'can_request_assets': can_request_assets(request.user),
        'highlight_id': request.GET.get('highlight', ''),
    })


# ─────────────────────────────────────────────────
# REQUEST FORM
# ─────────────────────────────────────────────────

@login_required
def mobile_request_form(request):
    requestable_items = filter_requestable_items(
        Item.objects.select_related('category').all().order_by('name'),
        request.user,
    )

    def _render_form(error=None, form_lines=None, reason='', notes='', is_urgent=False):
        return render(request, 'inventory/mobile/request_form.html', {
            'items': requestable_items,
            'error': error,
            'form_lines': form_lines or [{'item': request.GET.get('item', ''), 'quantity_requested': '1'}],
            'reason': reason,
            'notes': notes,
            'is_urgent': is_urgent,
            'preselected_item': request.GET.get('item', ''),
        })

    if request.method == 'POST':
        item_ids = request.POST.getlist('item')
        quantities = request.POST.getlist('quantity_requested')
        if not item_ids:
            item_ids = [request.POST.get('item', '')]
        if not quantities:
            quantities = [request.POST.get('quantity_requested', '1')]

        reason = request.POST.get('reason', '').strip()
        notes = request.POST.get('notes', '')
        is_urgent = request.POST.get('is_urgent') == 'on'
        form_lines = [
            {'item': item_id, 'quantity_requested': quantities[index] if index < len(quantities) else '1'}
            for index, item_id in enumerate(item_ids)
            if item_id or (index < len(quantities) and quantities[index])
        ]

        if not reason:
            return _render_form('Please provide a reason for your request.', form_lines, reason, notes, is_urgent)

        if not form_lines:
            return _render_form('Add at least one item.', form_lines, reason, notes, is_urgent)

        item_ids_clean = []
        request_lines = []
        for line in form_lines:
            item_id = line['item']
            if not item_id:
                return _render_form('Select an item for each request line.', form_lines, reason, notes, is_urgent)
            if item_id in item_ids_clean:
                return _render_form('Each item can only appear once in a request.', form_lines, reason, notes, is_urgent)
            item_ids_clean.append(item_id)
            try:
                quantity = int(line['quantity_requested'])
            except (ValueError, TypeError):
                return _render_form('Quantity must be a whole number.', form_lines, reason, notes, is_urgent)
            if quantity <= 0:
                return _render_form('Quantity must be at least 1.', form_lines, reason, notes, is_urgent)
            item = get_object_or_404(Item, id=item_id)
            if not user_can_request_item(request.user, item):
                return _render_form(UNAUTHORIZED_ITEM_REQUEST_MESSAGE, form_lines, reason, notes, is_urgent)
            if quantity > item.available_quantity:
                return _render_form(EXCESSIVE_QUANTITY_MESSAGE, form_lines, reason, notes, is_urgent)
            request_lines.append((item, quantity))

        with transaction.atomic():
            batch = ItemRequestBatch.objects.create(
                requested_by=request.user,
                purpose=reason,
                notes=notes,
                is_urgent=is_urgent,
            )
            ItemRequest.objects.bulk_create([
                ItemRequest(
                    batch=batch,
                    requested_by=request.user,
                    item=item,
                    quantity_requested=quantity,
                    units_remaining=item.available_quantity,
                    purpose=reason,
                    notes=notes,
                    is_urgent=is_urgent,
                    status='pending',
                )
                for item, quantity in request_lines
            ])
        return redirect('mobile_my_requests')

    return _render_form()


# ─────────────────────────────────────────────────
# MY REQUESTS
# ─────────────────────────────────────────────────

@login_required
def mobile_my_requests(request):
    requests_qs = ItemRequest.objects.filter(
        requested_by=request.user
    ).select_related('item').order_by('-requested_at')

    paginator = Paginator(requests_qs, 15)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'inventory/mobile/my_requests.html', {
        'requests': page_obj,
        'page_obj': page_obj,
    })


# ─────────────────────────────────────────────────
# MANAGER REQUESTS
# ─────────────────────────────────────────────────

@login_required
def mobile_manager_requests(request):
    user = request.user
    if getattr(user, 'role', '') not in ['admin', 'inventory_manager']:
        return redirect('mobile_dashboard')

    all_pending = ItemRequest.objects.filter(
        status='pending'
    ).select_related('requested_by', 'item', 'item__category').prefetch_related(
        'requisition_items__inventory_item'
    ).order_by('-requested_at')

    paginator = Paginator(all_pending, 5)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'inventory/mobile/manager_requests.html', {
        'requests': page_obj,
        'page_obj': page_obj,
        'total_pending': all_pending.count(),
    })


@login_required
def mobile_purchase_request_form(request):
    role = getattr(request.user, 'role', '')
    if role not in ['admin', 'inventory_manager']:
        return redirect('mobile_dashboard')

    categories = Category.objects.all().order_by('name')
    form_data = {
        'item_name': '',
        'item_type': 'consumable',
        'category': '',
        'quantity_requested': '1',
        'unit': 'pieces',
        'estimated_unit_cost': '',
        'needed_by': '',
        'urgency': 'normal',
        'reason': '',
    }

    if request.method == 'POST':
        form_data = {
            'item_name': (request.POST.get('item_name') or '').strip(),
            'item_type': (request.POST.get('item_type') or 'consumable').strip().lower(),
            'category': (request.POST.get('category') or '').strip(),
            'quantity_requested': (request.POST.get('quantity_requested') or '1').strip(),
            'unit': (request.POST.get('unit') or 'pieces').strip(),
            'estimated_unit_cost': (request.POST.get('estimated_unit_cost') or '').strip(),
            'needed_by': (request.POST.get('needed_by') or '').strip(),
            'urgency': (request.POST.get('urgency') or 'normal').strip(),
            'reason': (request.POST.get('reason') or '').strip(),
        }

        error = None
        category_obj = None

        if not form_data['item_name']:
            error = 'Please enter item/material name.'
        if not form_data['reason']:
            error = 'Please provide purchase reason.'

        try:
            quantity_requested = int(form_data['quantity_requested'])
            if quantity_requested < 1:
                raise ValueError
        except (TypeError, ValueError):
            error = error or 'Quantity must be at least 1.'
            quantity_requested = 1

        estimated_unit_cost = None
        if form_data['estimated_unit_cost']:
            try:
                estimated_unit_cost = Decimal(form_data['estimated_unit_cost'])
                if estimated_unit_cost < 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                error = error or 'Estimated unit cost must be a valid non-negative number.'

        if form_data['urgency'] not in ['normal', 'high']:
            form_data['urgency'] = 'normal'
        if form_data['item_type'] not in ['asset', 'consumable']:
            form_data['item_type'] = 'consumable'

        if form_data['category']:
            category_obj = Category.objects.filter(pk=form_data['category']).first()
            if not category_obj:
                error = error or 'Invalid category selected.'

        if error:
            return render(request, 'inventory/mobile/purchase_request_form.html', {
                'categories': categories,
                'error': error,
                'form': form_data,
            })

        PurchaseRequest.objects.create(
            requested_by=request.user,
            category=category_obj,
            item_name=form_data['item_name'],
            item_type=form_data['item_type'],
            quantity_requested=quantity_requested,
            unit=form_data['unit'] or 'pieces',
            estimated_unit_cost=estimated_unit_cost,
            needed_by=form_data['needed_by'] or None,
            urgency=form_data['urgency'],
            reason=form_data['reason'],
        )
        return redirect('/inventory/mobile/manager-requests/?pr=created')

    return render(request, 'inventory/mobile/purchase_request_form.html', {
        'categories': categories,
        'form': form_data,
    })


# ─────────────────────────────────────────────────
# ISSUED ITEMS
# ─────────────────────────────────────────────────

@login_required
def mobile_issued_items(request):
    role = getattr(request.user, 'role', '')
    is_manager = role in ['admin', 'inventory_manager']

    usage_logs_qs = ItemUsageLog.objects.select_related('logged_by').order_by('-logged_at')
    issues_qs = ItemIssue.objects.select_related(
        'request',
        'request__item',
        'request__item__category',
        'request__requested_by',
        'issued_by',
    ).prefetch_related(
        Prefetch('usage_logs', queryset=usage_logs_qs)
    ).annotate(
        total_consumed=Coalesce(
            Sum('usage_logs__quantity_used'),
            Value(0),
            output_field=DecimalField(max_digits=10, decimal_places=2),
        )
    ).order_by('-issued_at')

    if not is_manager:
        issues_qs = issues_qs.filter(request__requested_by=request.user)

    paginator = Paginator(issues_qs, 15)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    for issue in page_obj:
        total_used = issue.total_consumed or Decimal('0')
        issued_qty = Decimal(issue.issued_quantity)
        remaining = issued_qty - total_used
        issue.remaining_quantity = remaining if remaining > 0 else Decimal('0')
        if issued_qty > 0:
            issue.consumption_percent = int(min((total_used / issued_qty) * 100, 100))
        else:
            issue.consumption_percent = 0

    return render(request, 'inventory/mobile/issued_items.html', {
        'issues': page_obj,
        'page_obj': page_obj,
        'is_manager': is_manager,
    })


@login_required
def mobile_log_usage(request, pk):
    role = getattr(request.user, 'role', '')
    if role not in ['supervisor', 'caretaker']:
        return redirect('mobile_issued_items')

    if request.method != 'POST':
        return redirect('mobile_issued_items')

    issue = get_object_or_404(
        ItemIssue.objects.select_related(
            'request', 'request__item', 'request__requested_by'
        ),
        pk=pk,
        request__requested_by=request.user,
    )

    raw_qty = request.POST.get('quantity_used', '').strip()
    purpose = request.POST.get('purpose', '').strip()
    error = None
    quantity_used = None

    try:
        quantity_used = Decimal(raw_qty)
    except (ValueError, TypeError, InvalidOperation):
        error = 'Enter a valid number.'

    if error is None and quantity_used is not None:
        if quantity_used <= 0:
            error = 'Quantity must be greater than 0.'
        else:
            quantity_used = quantity_used.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            if issue.request.item.item_type == 'asset':
                if quantity_used != quantity_used.to_integral_value():
                    error = 'Enter a whole number for assets.'

    if error is None and not purpose:
        error = 'Please describe what the item was used for.'

    if error is None:
        with transaction.atomic():
            issue = (
                ItemIssue.objects
                .select_for_update()
                .select_related('request__item', 'request__requested_by')
                .get(pk=pk, request__requested_by=request.user)
            )
            total_used = ItemUsageLog.objects.filter(issue=issue).aggregate(
                total=Coalesce(
                    Sum('quantity_used'),
                    Value(0),
                    output_field=DecimalField(max_digits=10, decimal_places=2),
                )
            )['total']
            remaining = Decimal(issue.issued_quantity) - total_used
            if remaining <= 0:
                error = 'All issued quantity has already been logged.'
            elif quantity_used > remaining:
                remaining_display = remaining.quantize(Decimal('0.01'))
                error = (
                    f'Cannot exceed remaining quantity ({remaining_display} '
                    f'{issue.request.item.unit}).'
                )
            else:
                ItemUsageLog.objects.create(
                    issue=issue,
                    quantity_used=quantity_used,
                    purpose=purpose,
                    logged_by=request.user,
                )
        if error is None:
            return redirect('mobile_issued_items')

    usage_logs_qs = ItemUsageLog.objects.select_related('logged_by').order_by('-logged_at')
    issues_qs = ItemIssue.objects.select_related(
        'request', 'request__item', 'request__item__category',
        'request__requested_by', 'issued_by',
    ).prefetch_related(
        Prefetch('usage_logs', queryset=usage_logs_qs)
    ).annotate(
        total_consumed=Coalesce(
            Sum('usage_logs__quantity_used'),
            Value(0),
            output_field=DecimalField(max_digits=10, decimal_places=2),
        )
    ).filter(request__requested_by=request.user).order_by('-issued_at')

    paginator = Paginator(issues_qs, 15)
    page_obj = paginator.get_page(1)

    for issue in page_obj:
        total_used = issue.total_consumed or Decimal('0')
        issued_qty = Decimal(issue.issued_quantity)
        remaining = issued_qty - total_used
        issue.remaining_quantity = remaining if remaining > 0 else Decimal('0')
        if issued_qty > 0:
            issue.consumption_percent = int(min((total_used / issued_qty) * 100, 100))
        else:
            issue.consumption_percent = 0

    return render(request, 'inventory/mobile/issued_items.html', {
        'issues': page_obj,
        'page_obj': page_obj,
        'is_manager': False,
        'form_error': error,
        'error_issue_id': issue.id,
        'open_form_issue_id': issue.id,
    })


# ─────────────────────────────────────────────────
# STOCK LOG
# ─────────────────────────────────────────────────

@login_required
def mobile_stock_log(request):
    role = getattr(request.user, 'role', '')
    if role not in ['admin', 'inventory_manager']:
        return redirect('mobile_dashboard')

    entries = StockEntry.objects.select_related('item', 'item__category', 'added_by').order_by('-added_at')
    paginator = Paginator(entries, 20)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'inventory/mobile/stock_log.html', {
        'entries': page_obj,
        'page_obj': page_obj,
    })


# ─────────────────────────────────────────────────
# REPORTS
# ─────────────────────────────────────────────────

@login_required
def mobile_reports(request):
    role = getattr(request.user, 'role', '')
    if role not in ['admin', 'inventory_manager']:
        return redirect('mobile_dashboard')

    from apps.reports.generators import generate_inventory_report, generate_requests_report

    report_type = request.GET.get('type', 'inventory')
    status_filter = request.GET.get('filter', 'all')
    item_type_filter = (request.GET.get('item_type') or 'all').strip().lower()
    if item_type_filter not in {'all', 'asset', 'consumable'}:
        item_type_filter = 'all'
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    page_number = request.GET.get('page', 1)

    inventory_data = None
    requests_data = None
    inventory_page_obj = None
    requests_page_obj = None

    if report_type == 'requests':
        requests_data = generate_requests_report(date_from or None, date_to or None)
        requests_rows = requests_data.get('requests', [])
        requests_page_obj = Paginator(requests_rows, 15).get_page(page_number)
    else:
        inventory_data = generate_inventory_report(
            status_filter=status_filter if status_filter and status_filter != 'all' else None,
            item_type_filter=item_type_filter if item_type_filter != 'all' else None,
        )
        inventory_rows = inventory_data.get('items', [])
        inventory_page_obj = Paginator(inventory_rows, 15).get_page(page_number)

    return render(request, 'inventory/mobile/reports.html', {
        'report_type': report_type,
        'inventory_data': inventory_data,
        'inventory_page_obj': inventory_page_obj,
        'requests_data': requests_data,
        'requests_page_obj': requests_page_obj,
        'status_filter': status_filter or 'all',
        'item_type_filter': item_type_filter,
        'date_from': date_from,
        'date_to': date_to,
    })


def _export_report_multiformat(export_format, headers, rows, filename_base, title):
    fmt = (export_format or 'csv').strip().lower()

    if fmt == 'xlsx':
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active if wb.active is not None else wb.create_sheet(title=title[:30])
            ws.title = title[:30]

            ws.append([title])
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
            title_cell = ws.cell(row=1, column=1)
            title_cell.font = Font(name='Arial', size=13, bold=True, color='FFFFFF')
            title_cell.fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
            title_cell.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 28

            ws.append(headers)
            ws.row_dimensions[2].height = 22
            header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
            header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=2, column=col_idx)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.border = border

            alt_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
            for row_idx, row in enumerate(rows, start=3):
                ws.append(row)
                ws.row_dimensions[row_idx].height = 18
                is_alt = (row_idx % 2 == 0)
                for col_idx in range(1, len(row) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = Font(name='Arial', size=9)
                    cell.border = border
                    if is_alt:
                        cell.fill = alt_fill

            for col_idx, col in enumerate(ws.columns, 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            response = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="{filename_base}.xlsx"'
            return response
        except Exception:
            pass

    if fmt == 'pdf':
        try:
            from reportlab.lib.pagesizes import letter, landscape
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors

            buf = io.BytesIO()
            doc = SimpleDocTemplate(
                buf,
                pagesize=landscape(letter) if len(headers) > 6 else letter,
                rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24
            )
            elements: list = []
            styles = getSampleStyleSheet()

            elements.append(Paragraph(f"<b>{title}</b>", styles['Title']))
            elements.append(Spacer(1, 10))

            table_data = [headers] + [[str(c if c is not None else '') for c in r] for r in rows]
            table = Table(table_data, repeatRows=1)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 8),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 5),
                ('TOPPADDING', (0, 0), (-1, 0), 5),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 7),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ]))
            elements.append(table)
            doc.build(elements)

            buf.seek(0)
            response = HttpResponse(buf.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="{filename_base}.pdf"'
            return response
        except Exception:
            pass

    if fmt == 'print':
        html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{title}</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; margin: 20px; color: #1e293b; }}
        h1 {{ font-size: 18px; margin-bottom: 4px; color: #0f172a; }}
        .meta {{ font-size: 11px; color: #64748b; margin-bottom: 16px; }}
        table {{ width: 100%; border-collapse: collapse; font-size: 11px; }}
        th, td {{ border: 1px solid #cbd5e1; padding: 6px 8px; text-align: left; }}
        th {{ background: #f1f5f9; font-weight: 700; }}
        tr:nth-child(even) {{ background: #f8fafc; }}
        @media print {{
            @page {{ size: landscape; margin: 12mm; }}
            body {{ margin: 0; }}
        }}
    </style>
</head>
<body>
    <h1>{title}</h1>
    <div class="meta">Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}</div>
    <table>
        <thead>
            <tr>{''.join(f'<th>{h}</th>' for h in headers)}</tr>
        </thead>
        <tbody>
            {''.join('<tr>' + ''.join(f'<td>{c if c is not None else ""}</td>' for c in row) + '</tr>' for row in rows)}
        </tbody>
    </table>
    <script>window.onload = function() {{ window.print(); }};</script>
</body>
</html>"""
        return HttpResponse(html, content_type='text/html')

    # Default: CSV
    filename = f'{filename_base}.csv'
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write('\ufeff')

    writer = csv.writer(response)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)

    return response


@login_required
def mobile_export_inventory_report_csv(request):
    role = getattr(request.user, 'role', '')
    if role not in ['admin', 'inventory_manager']:
        return redirect('mobile_dashboard')

    from apps.reports.generators import generate_inventory_report

    status_filter = (request.GET.get('filter') or 'all').strip().lower()
    item_type_filter = (request.GET.get('item_type') or 'all').strip().lower()
    export_format = (request.GET.get('format') or 'csv').strip().lower()

    if status_filter not in {'all', 'in_stock', 'low_stock', 'out_of_stock'}:
        status_filter = 'all'
    if item_type_filter not in {'all', 'asset', 'consumable'}:
        item_type_filter = 'all'

    report = generate_inventory_report(
        status_filter=status_filter if status_filter != 'all' else None,
        item_type_filter=item_type_filter if item_type_filter != 'all' else None,
    )

    filename_base = f'inventory_report_{status_filter}_{item_type_filter}'
    headers = [
        'ID', 'Name', 'Category', 'Type', 'Unit',
        'Total Quantity', 'Available Quantity', 'Threshold Quantity',
        'Unit Cost', 'Total Value', 'Status',
    ]
    rows = []
    for row in report.get('items', []):
        rows.append([
            row.get('id'),
            row.get('name'),
            row.get('category'),
            row.get('type'),
            row.get('unit'),
            row.get('total_quantity'),
            row.get('available_quantity'),
            row.get('threshold_quantity'),
            row.get('unit_cost'),
            row.get('total_value'),
            row.get('status'),
        ])

    return _export_report_multiformat(
        export_format,
        headers,
        rows,
        filename_base,
        f'Inventory Status Report ({status_filter.replace("_", " ").title()})',
    )


@login_required
def mobile_export_requests_report_csv(request):
    role = getattr(request.user, 'role', '')
    if role not in ['admin', 'inventory_manager']:
        return redirect('mobile_dashboard')

    from apps.reports.generators import generate_requests_report

    date_from_raw = (request.GET.get('date_from') or '').strip()
    date_to_raw = (request.GET.get('date_to') or '').strip()
    export_format = (request.GET.get('format') or 'csv').strip().lower()
    date_from = parse_date(date_from_raw)
    date_to = parse_date(date_to_raw)

    report = generate_requests_report(
        date_from.isoformat() if date_from else None,
        date_to.isoformat() if date_to else None,
    )

    filename_base = 'requests_report'
    headers = [
        'ID', 'Requester', 'Item', 'Quantity', 'Status', 'Urgent',
        'Reason', 'Created At', 'Reviewed By', 'Reviewed At', 'Issued At',
    ]
    rows = []
    for row in report.get('requests', []):
        rows.append([
            row.get('id'),
            row.get('requester'),
            row.get('item'),
            row.get('quantity'),
            row.get('status'),
            'Yes' if row.get('is_urgent') else 'No',
            row.get('reason'),
            row.get('created_at'),
            row.get('reviewed_by') or '',
            row.get('reviewed_at') or '',
            row.get('issued_at') or '',
        ])

    return _export_report_multiformat(
        export_format,
        headers,
        rows,
        filename_base,
        'Inventory Requests Report',
    )


# ─────────────────────────────────────────────────
# PROFILE
# ─────────────────────────────────────────────────

@login_required
def mobile_profile(request):
    return render(request, 'inventory/mobile/profile.html', {
        'user': request.user,
    })


@login_required
def mobile_change_password(request):
    if request.method == 'POST':
        current = request.POST.get('current_password')
        new_pass = request.POST.get('new_password')
        confirm = request.POST.get('confirm_password')
        error = None

        if not request.user.check_password(current):
            error = 'Current password is incorrect.'
        elif new_pass != confirm:
            error = 'New passwords do not match.'
        elif len(new_pass) < 8:
            error = 'Password must be at least 8 characters.'

        if error:
            return render(request, 'inventory/mobile/profile.html', {
                'user': request.user,
                'error': error,
            })

        request.user.set_password(new_pass)
        request.user.save()
        update_session_auth_hash(request, request.user)
        return render(request, 'inventory/mobile/profile.html', {
            'user': request.user,
            'success': 'Password changed successfully.',
        })

    return redirect('mobile_profile')


# ─────────────────────────────────────────────────
# REQUEST ACTIONS
# ─────────────────────────────────────────────────

@login_required
@transaction.atomic
def mobile_approve_request(request, pk):
    """Approve a pending request (manager/admin only)."""
    if request.method == 'POST':
        role = getattr(request.user, 'role', '')
        if role not in ['admin', 'inventory_manager']:
            return redirect('mobile_dashboard')

        item_request = get_object_or_404(
            ItemRequest.objects.select_for_update().prefetch_related('requisition_items'), pk=pk
        )
        if item_request.status == 'pending':
            lines = item_request.requisition_lines()
            item_ids = [line.inventory_item_id for line in lines]
            locked_items = {
                item.id: item
                for item in Item.objects.select_for_update().filter(pk__in=item_ids)
            }
            for line in lines:
                item = locked_items.get(line.inventory_item_id)
                available_to_reserve = (
                    item.available_quantity - item.reserved_quantity
                    if item else 0
                )
                if not item or line.quantity > available_to_reserve:
                    return redirect('mobile_manager_requests')
            for line in lines:
                Item.objects.filter(pk=line.inventory_item_id).update(
                    reserved_quantity=F('reserved_quantity') + line.quantity
                )
            first_line = next((line for line in lines if line.inventory_item_id == item_request.item_id), None)
            first_item = locked_items.get(item_request.item_id)
            item_request.units_remaining = max(
                (
                    (first_item.available_quantity - first_item.reserved_quantity)
                    if first_item else 0
                ) - (first_line.quantity if first_line else 0),
                0,
            )
            item_request.status = 'approved'
            item_request.approved_by = request.user
            item_request.reviewed_at = timezone.now()
            item_request.save(update_fields=[
                'status', 'approved_by', 'reviewed_at', 'units_remaining'
            ])
            try:
                from apps.housekeeping.models import HKStockRequest
                hk_req = HKStockRequest.objects.filter(inventory_request=item_request).first()
                if hk_req:
                    hk_req.status = 'approved'
                    hk_req.save(update_fields=['status'])
            except Exception:
                pass
            try:
                from apps.horticulture.models import HorticultureStockRequest
                horti_req = HorticultureStockRequest.objects.filter(inventory_request=item_request).first()
                if horti_req:
                    horti_req.status = 'approved'
                    horti_req.save(update_fields=['status'])
            except Exception:
                pass
    return redirect('mobile_manager_requests')


@login_required
def mobile_reject_request(request, pk):
    """Reject a pending request with an optional reason."""
    if request.method == 'POST':
        role = getattr(request.user, 'role', '')
        if role not in ['admin', 'inventory_manager']:
            return redirect('mobile_dashboard')

        item_request = get_object_or_404(ItemRequest, pk=pk)
        if item_request.status == 'pending':
            item_request.status = 'rejected'
            item_request.approved_by = request.user
            item_request.rejection_reason = request.POST.get('rejection_reason', '').strip()
            item_request.reviewed_at = timezone.now()
            item_request.units_remaining = item_request.item.available_quantity
            item_request.save(update_fields=[
                'status', 'approved_by', 'rejection_reason', 'reviewed_at', 'units_remaining'
            ])
            try:
                from apps.housekeeping.models import HKStockRequest
                hk_req = HKStockRequest.objects.filter(inventory_request=item_request).first()
                if hk_req:
                    hk_req.status = 'rejected'
                    hk_req.save(update_fields=['status'])
            except Exception:
                pass
            try:
                from apps.horticulture.models import HorticultureStockRequest
                horti_req = HorticultureStockRequest.objects.filter(inventory_request=item_request).first()
                if horti_req:
                    horti_req.status = 'rejected'
                    horti_req.save(update_fields=['status'])
            except Exception:
                pass
    return redirect('mobile_manager_requests')


@login_required
@transaction.atomic
def mobile_issue_request(request, pk):
    """Issue an approved request: deduct stock and create ItemIssue record."""
    if request.method == 'POST':
        role = getattr(request.user, 'role', '')
        if role not in ['admin', 'inventory_manager']:
            return redirect('mobile_dashboard')

        item_request = get_object_or_404(
            ItemRequest.objects.select_related('item').select_for_update(), pk=pk
        )

        # Allow issuing directly from pending (auto-approve + issue in one step)
        if item_request.status not in ['pending', 'approved']:
            return redirect('mobile_manager_requests')

        item = item_request.item

        try:
            qty = int(request.POST.get('quantity', item_request.quantity_requested))
        except (ValueError, TypeError):
            qty = item_request.quantity_requested

        qty = max(1, qty)
        expected_return_date = None

        available_to_issue = item.available_quantity
        if (
            item_request.status == 'approved'
            and item.reserved_quantity >= item_request.quantity_requested
        ):
            available_to_issue = min(item.available_quantity, item.reserved_quantity)
        if available_to_issue < qty:
            # Not enough stock — redirect back without issuing
            return redirect('mobile_manager_requests')

        if item.item_type == Item.ItemType.ASSET:
            raw_return_date = (request.POST.get('expected_return_date') or '').strip()
            if raw_return_date:
                expected_return_date = parse_date(raw_return_date)
                if expected_return_date is None or expected_return_date < timezone.localdate():
                    return redirect('mobile_manager_requests')

        reserved_release = min(item.reserved_quantity, item_request.quantity_requested)
        Item.objects.filter(pk=item.pk).update(
            available_quantity=F('available_quantity') - qty,
            reserved_quantity=F('reserved_quantity') - reserved_release,
        )

        # If still pending, mark it approved first
        if item_request.status == 'pending':
            item_request.approved_by = request.user
            item_request.reviewed_at = timezone.now()
            item_request.units_remaining = max(item.available_quantity - qty, 0)

        item_request.status = 'issued'
        item_request.save()

        ItemIssue.objects.update_or_create(
            request=item_request,
            defaults={
                'issued_by': request.user,
                'issued_quantity': qty,
                'expected_return_date': expected_return_date,
                'notes': request.POST.get('notes', ''),
            },
        )
        try:
            from apps.housekeeping.models import HKStockRequest
            hk_req = HKStockRequest.objects.filter(inventory_request=item_request).first()
            if hk_req:
                hk_req.status = 'issued'
                hk_req.save(update_fields=['status'])
        except Exception:
            pass
        try:
            from apps.horticulture.models import HorticultureStockRequest
            horti_req = HorticultureStockRequest.objects.filter(inventory_request=item_request).first()
            if horti_req:
                horti_req.status = 'issued'
                horti_req.save(update_fields=['status'])
        except Exception:
            pass
    return redirect('mobile_manager_requests')
